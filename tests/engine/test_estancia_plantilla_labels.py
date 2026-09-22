"""Fix reglas #9/#33/#44/#61/#70/#74: plantilla Estancia + live overrides.

- Opt-in plantilla 'Estancia: {estancia_str}' en group path.
- infer_aggregations infiere compute_horas para hospi_sala_obs_cantidad_check.
- Live override cups / tipo-edad en normalized_rows.
"""

from __future__ import annotations

from datetime import datetime, timedelta


def _ce():
    from app.services.engine.condition_evaluator import ConditionEvaluator
    return ConditionEvaluator()


def _collector():
    from app.services.engine.evidence_collector import EvidenceCollector
    return EvidenceCollector()


def _always_true_horas_tree():
    return _ce().build_tree([
        {"id": 1, "padre_id": None, "tipo": "atomic", "operador": "gt",
         "fuente_datos": "group.estancia_horas", "valor_esperado": 1, "orden": 0},
    ])


def _mk_group_ws(fec0: datetime, fec1: datetime):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="NUMERO_FACTURA")
    ws.cell(row=1, column=2, value="FEC_FACTURA")
    ws.cell(row=1, column=3, value="FECHA_CIERRE")
    ws.cell(row=2, column=1, value="F001")
    ws.cell(row=2, column=2, value=fec0)
    ws.cell(row=2, column=3, value=fec1)
    indices = {"numero_factura": 0, "fec_factura": 1, "fecha_cierre": 2}
    return ws, indices


def _agg_horas():
    return [
        {"function": "compute_horas", "field1": "fec_factura",
         "field2": "fecha_cierre", "target": "estancia_horas"},
    ]


class TestPlantillaEstanciaOptIn:
    def test_detalle_b_plantilla_con_label_enriquece(self):
        from app.services.engine.group_evaluator import GroupEvaluator
        base = datetime(2026, 1, 1, 8, 0, 0)
        ws, indices = _mk_group_ws(base, base + timedelta(hours=30, minutes=30))
        groups = GroupEvaluator.build_groups(ws, indices)
        rule_info = {
            "id": 70, "version": 1, "dominio": "urgencias",
            "nombre": "hospi_sala", "descripcion": "Estancia",
            "severidad": "error",
            "detalle_b_campo": "Estancia: {estancia_str}",
        }
        results = GroupEvaluator.evaluate(
            groups, ws, indices, _agg_horas(), _always_true_horas_tree(),
            _ce(), rule_info, _collector(), record_evidence=False,
        )
        assert len(results) == 1
        assert results[0].get("estancia_str") == "1d 6h"


class TestInferHospi:
    def test_hospi_operador_infiere_compute_horas(self):
        from app.services.engine.group_evaluator import GroupEvaluator
        tree = _ce().build_tree([
            {"id": 1, "padre_id": None, "tipo": "atomic",
             "operador": "hospi_sala_obs_cantidad_check",
             "fuente_datos": "invoice.codigo",
             "valor_esperado": "38114", "orden": 0},
        ])
        aggs = GroupEvaluator.infer_aggregations(tree)
        assert any(a.get("function") == "compute_horas" for a in aggs)


class TestLiveOverrideCups:
    def test_cups_live_override_pisa_formatter(self):
        from app.services.normalized_rows import build_normalized_rows
        item = {
            "factura": "F001", "codigo": "C001", "procedimiento": "PROC",
            "estancia_str": "5h", "problema": "P",
            "detalle_a_campo": "codigo",
            "detalle_b_campo": "Estancia: {estancia_str}",
        }
        rows = build_normalized_rows(
            error_groups={"Cups-Equivalentes": [item]},
            responsables_map={}, fec_factura_map={}, fecha_cierre_vacia_map={},
            grupo_mappings={"Cups-Equivalentes": {}},
        )
        assert len(rows) == 1
        assert rows[0]["procedimiento"] == "C001"
        assert rows[0]["detalle"] == "Estancia: 5h"


class TestLiveOverrideTipoEdad:
    def test_tipo_edad_live_override_pisa_formatter(self):
        from app.services.normalized_rows import build_normalized_rows
        item = {
            "factura": "F001", "identificacion": "123",
            "fec_nacimiento": "2000-01-01", "fec_factura": "2026-01-01",
            "tipo_actual": "TI", "tipo_deberia": "CC",
            "problema": "Tipo mal",
            "date.edad": 26,
            "detalle_a_campo": "identificacion",
            "detalle_b_campo": "Edad: {date.edad}",
        }
        rows = build_normalized_rows(
            error_groups={"Tipo Identificacion / Edad": [item]},
            responsables_map={}, fec_factura_map={}, fecha_cierre_vacia_map={},
            grupo_mappings={"Tipo Identificacion / Edad": {}},
        )
        assert len(rows) == 1
        assert rows[0]["procedimiento"] == "123"
        assert rows[0]["detalle"] == "Edad: 26"
