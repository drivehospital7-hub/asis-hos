"""Engine enriquece problem dict con estancia_str (formato legacy).

TDD RED: estos tests deben FALLAR antes del fix porque el engine pone
estancia_horas pero jamas estancia_str, y detalle_b_campo='estancia_str'
(regla #44) resuelve vacio.
"""

from __future__ import annotations

from datetime import datetime, timedelta
from unittest.mock import MagicMock


def _ce():
    from app.services.engine.condition_evaluator import ConditionEvaluator
    return ConditionEvaluator()


def _collector():
    from app.services.engine.evidence_collector import EvidenceCollector
    return EvidenceCollector()


def _rule_info(**over):
    base = {
        "id": 44, "version": 1, "dominio": "urgencias",
        "nombre": "sala_obs_estancia", "descripcion": "Estancia test",
        "severidad": "error", "detalle_b_campo": "estancia_str",
    }
    base.update(over)
    return base


def _always_true_tree():
    return _ce().build_tree([
        {"id": 1, "padre_id": None, "tipo": "atomic", "operador": "gt",
         "fuente_datos": "group.estancia_horas", "valor_esperado": 1, "orden": 0},
    ])


def _mk_group_ws(fec0: datetime | None, fec1: datetime | None):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="NUMERO_FACTURA")
    ws.cell(row=1, column=2, value="FEC_FACTURA")
    ws.cell(row=1, column=3, value="FECHA_CIERRE")
    if fec0 is not None:
        ws.cell(row=2, column=1, value="F001")
        ws.cell(row=2, column=2, value=fec0)
        ws.cell(row=2, column=3, value=fec1)
    indices = {"numero_factura": 0, "fec_factura": 1, "fecha_cierre": 2}
    return ws, indices


def _agg_horas():
    return [
        {"function": "compute_horas", "field1": "fec_factura",
         "field2": "fecha_cierre", "target": "estancia_horas"},
    ]


class TestGroupEstanciaStr:
    def test_30_5_horas_formatea_1d_6h(self):
        from app.services.engine.group_evaluator import GroupEvaluator
        base = datetime(2026, 1, 1, 8, 0, 0)
        ws, indices = _mk_group_ws(base, base + timedelta(hours=30, minutes=30))
        groups = GroupEvaluator.build_groups(ws, indices)
        results = GroupEvaluator.evaluate(
            groups, ws, indices, _agg_horas(), _always_true_tree(),
            _ce(), _rule_info(), _collector(), record_evidence=False,
        )
        assert len(results) == 1
        assert results[0].get("estancia_horas") == 30.5
        assert results[0].get("estancia_str") == "1d 6h"

    def test_5_horas_formatea_5h(self):
        from app.services.engine.group_evaluator import GroupEvaluator
        base = datetime(2026, 1, 1, 8, 0, 0)
        ws, indices = _mk_group_ws(base, base + timedelta(hours=5))
        groups = GroupEvaluator.build_groups(ws, indices)
        results = GroupEvaluator.evaluate(
            groups, ws, indices, _agg_horas(), _always_true_tree(),
            _ce(), _rule_info(), _collector(), record_evidence=False,
        )
        assert len(results) == 1
        assert results[0].get("estancia_str") == "5h"

    def test_sin_estancia_sin_clave(self):
        """Sin agregacion de horas no hay estancia_horas -> sin estancia_str (fallback intacto)."""
        from app.services.engine.group_evaluator import GroupEvaluator
        from openpyxl import Workbook
        wb = Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="NUMERO_FACTURA")
        ws.cell(row=1, column=2, value="CODIGO")
        ws.cell(row=2, column=1, value="F001")
        ws.cell(row=2, column=2, value="5DSB01")
        indices = {"numero_factura": 0, "codigo": 1}
        groups = GroupEvaluator.build_groups(ws, indices)
        flat = _ce().build_tree([
            {"id": 1, "padre_id": None, "tipo": "atomic", "operador": "eq",
             "fuente_datos": "invoice.codigo", "valor_esperado": "5DSB01", "orden": 0},
        ])
        results = GroupEvaluator.evaluate(
            groups, ws, indices, [], flat,
            _ce(), _rule_info(), _collector(), record_evidence=False,
        )
        assert len(results) == 1
        assert "estancia_str" not in results[0]

    def test_no_sobrescribe_estancia_str_existente(self):
        from app.services.engine.group_evaluator import GroupEvaluator
        base = datetime(2026, 1, 1, 8, 0, 0)
        ws, indices = _mk_group_ws(base, base + timedelta(hours=30, minutes=30))
        groups = GroupEvaluator.build_groups(ws, indices)
        # Pre-existente a nivel de fila: el bridge lo propaga y no se pisa.
        ws.cell(row=2, column=1, value="F001")
        results = GroupEvaluator.evaluate(
            groups, ws, indices, _agg_horas(), _always_true_tree(),
            _ce(), _rule_info(), _collector(), record_evidence=False,
            rows=[{"numero_factura": "F001", "fec_factura": base,
                   "fecha_cierre": base + timedelta(hours=30, minutes=30),
                   "estancia_str": "custom"}],
        )
        assert len(results) == 1
        assert results[0].get("estancia_str") == "custom"


def _make_cond(**over):
    m = MagicMock()
    for k, v in over.items():
        setattr(m, k, v)
    return m


class TestEngineRowEstanciaStr:
    def test_row_path_enriquece_desde_estancia_horas(self):
        """Fila con estancia_horas=30.5 -> problem trae estancia_str='1d 6h'."""
        from app.services.engine.engine import RuleEvaluationEngine
        from app.models import Regla

        rule = Regla(
            id=44, nombre="sala_obs_estancia", dominio="urgencias",
            estado="active", version=1, prioridad=10, severidad="error",
        )
        rule.detalle_b_campo = "estancia_str"
        root = {
            "id": 1, "regla_id": 44, "padre_id": None,
            "tipo": "atomic", "operador": "eq",
            "fuente_datos": "invoice.codigo",
            "valor_esperado": "5DSB01", "orden": 0,
        }
        session = MagicMock()
        q = MagicMock()
        q.filter.return_value = q
        q.order_by.return_value = q
        q.first.return_value = rule
        q.all.return_value = [_make_cond(**root)]
        session.query.return_value = q

        engine = RuleEvaluationEngine(session)
        rows = [{"numero_factura": "F001", "codigo": "5DSB01",
                 "estancia_horas": 30.5}]
        results = engine.evaluate_sheet(
            "sala_obs_estancia", None,
            {"numero_factura": 0, "codigo": 1},
            persist=False, rows=rows,
        )
        assert len(results) == 1
        assert results[0].get("estancia_str") == "1d 6h"

    def test_row_path_5h(self):
        from app.services.engine.engine import RuleEvaluationEngine
        from app.models import Regla

        rule = Regla(
            id=44, nombre="sala_obs_estancia", dominio="urgencias",
            estado="active", version=1, prioridad=10, severidad="error",
        )
        rule.detalle_b_campo = "estancia_str"
        root = {
            "id": 1, "regla_id": 44, "padre_id": None,
            "tipo": "atomic", "operador": "eq",
            "fuente_datos": "invoice.codigo",
            "valor_esperado": "5DSB01", "orden": 0,
        }
        session = MagicMock()
        q = MagicMock()
        q.filter.return_value = q
        q.order_by.return_value = q
        q.first.return_value = rule
        q.all.return_value = [_make_cond(**root)]
        session.query.return_value = q

        engine = RuleEvaluationEngine(session)
        rows = [{"numero_factura": "F001", "codigo": "5DSB01",
                 "estancia_horas": 5.0}]
        results = engine.evaluate_sheet(
            "sala_obs_estancia", None,
            {"numero_factura": 0, "codigo": 1},
            persist=False, rows=rows,
        )
        assert len(results) == 1
        assert results[0].get("estancia_str") == "5h"

    def test_row_path_sin_estancia_sin_clave(self):
        from app.services.engine.engine import RuleEvaluationEngine
        from app.models import Regla

        rule = Regla(
            id=44, nombre="sala_obs_estancia", dominio="urgencias",
            estado="active", version=1, prioridad=10, severidad="error",
        )
        rule.detalle_b_campo = "estancia_str"
        root = {
            "id": 1, "regla_id": 44, "padre_id": None,
            "tipo": "atomic", "operador": "eq",
            "fuente_datos": "invoice.codigo",
            "valor_esperado": "5DSB01", "orden": 0,
        }
        session = MagicMock()
        q = MagicMock()
        q.filter.return_value = q
        q.order_by.return_value = q
        q.first.return_value = rule
        q.all.return_value = [_make_cond(**root)]
        session.query.return_value = q

        engine = RuleEvaluationEngine(session)
        rows = [{"numero_factura": "F001", "codigo": "5DSB01"}]
        results = engine.evaluate_sheet(
            "sala_obs_estancia", None,
            {"numero_factura": 0, "codigo": 1},
            persist=False, rows=rows,
        )
        assert len(results) == 1
        assert "estancia_str" not in results[0]
