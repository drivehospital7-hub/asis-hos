"""Paridad fila ↔ grupo sin migrar condiciones (puente automático).

Contrato: una condición escrita para fila debe dar el mismo veredicto en
modo grupo cuando el grupo es homogéneo. El modo grupo resuelve por dentro:

- ``date.horas`` desde el primer par válido (fec_factura, fecha_cierre).
- ``invoice.codigo`` con ``in``/``set_*`` contra el collect_set del grupo
  (computado on-demand si el param no trae la aggregation).
- escalares (tipo_factura_descripcion, ...) desde la primera fila.

Caso testigo #44 sala_obs_mayor_6_horas:
AND(gte(date.horas, 6), eq(tipo_factura, Urgencias), NOT(in(codigo, SALA))).
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook

SALA_CODES = ["5DSB01", "05DSB01", "129B02", "38114", "38915"]


def _conds_44() -> list[dict]:
    """Condición de fila #44 escrita UNA sola vez (sin fuentes de grupo)."""
    return [
        {"id": 1, "padre_id": None, "tipo": "composite", "operador": "AND",
         "fuente_datos": None, "valor_esperado": None, "orden": 0},
        {"id": 2, "padre_id": 1, "tipo": "atomic", "operador": "gte",
         "fuente_datos": "date.horas", "valor_esperado": 6, "orden": 0},
        {"id": 3, "padre_id": 1, "tipo": "atomic", "operador": "eq",
         "fuente_datos": "invoice.tipo_factura_descripcion",
         "valor_esperado": "Urgencias", "orden": 1},
        {"id": 4, "padre_id": 1, "tipo": "composite", "operador": "NOT",
         "fuente_datos": None, "valor_esperado": None, "orden": 2},
        {"id": 5, "padre_id": 4, "tipo": "atomic", "operador": "in",
         "fuente_datos": "invoice.codigo", "valor_esperado": SALA_CODES,
         "orden": 0},
    ]


def _row_verdict(row: dict) -> bool:
    """Veredicto modo fila para una sola fila."""
    from app.services.engine.condition_evaluator import ConditionEvaluator
    from app.services.engine.context import EvaluationContext

    ce = ConditionEvaluator()
    tree = ce.build_tree(_conds_44())
    ctx = EvaluationContext(invoice_data=dict(row))
    return bool(ce.evaluate(tree, ctx).get("outcome", False))


def _group_results(rows_list: list[dict], agg_configs: list | None = None):
    """Veredicto modo grupo (UNA detección por factura como máximo)."""
    from app.services.engine.condition_evaluator import ConditionEvaluator
    from app.services.engine.evidence_collector import EvidenceCollector
    from app.services.engine.group_evaluator import GroupEvaluator

    ce = ConditionEvaluator()
    tree = ce.build_tree(_conds_44())
    groups = GroupEvaluator.build_groups(rows=rows_list)
    rule_info = {"id": 44, "version": 1, "dominio": "urgencias",
                 "nombre": "sala_obs_mayor_6_horas",
                 "descripcion": "Sala obs > 6h", "severidad": "error"}
    return GroupEvaluator.evaluate(
        groups, None, {}, agg_configs or [], tree, ce, rule_info,
        EvidenceCollector(), record_evidence=False, rows=rows_list,
    )


def _row_7h(codigo: str = "890701", factura: str = "F001",
            cierre: str = "2026-06-24 15:00:00") -> dict:
    return {"numero_factura": factura, "codigo": codigo,
            "tipo_factura_descripcion": "Urgencias",
            "fec_factura": "2026-06-24 08:00:00", "fecha_cierre": cierre}


class TestParidad44:
    def test_7h_sin_sala_match_una_sola_vez(self):
        rows_list = [_row_7h("890701")]
        assert _row_verdict(rows_list[0]) is True
        results = _group_results(rows_list)
        assert len(results) == 1
        assert results[0]["factura"] == "F001"

    def test_con_codigo_sala_no_match(self):
        rows_list = [_row_7h("5DSB01")]
        assert _row_verdict(rows_list[0]) is False
        assert _group_results(rows_list) == []

    def test_3h_no_match(self):
        rows_list = [_row_7h("890701",
                             cierre="2026-06-24 11:00:00")]
        assert _row_verdict(rows_list[0]) is False
        assert _group_results(rows_list) == []

    def test_grupo_homogeneo_multifila_mismo_veredicto(self):
        rows_list = [_row_7h("890701"), _row_7h("890701")]
        assert all(_row_verdict(r) for r in rows_list)
        results = _group_results(rows_list)
        assert len(results) == 1
        assert results[0]["factura"] == "F001"

    def test_grupo_heterogeneo_divergencia_documentada(self):
        """Filas difieren entre sí; el grupo da UN solo veredicto.

        Semántica de grupo para ``in``: hay MATCH si ALGÚN código del
        grupo está en el set (equivale a set_intersects). Con NOT, el
        grupo solo hace MATCH si NINGÚN código es de sala.
        """
        rows_list = [_row_7h("890701"), _row_7h("5DSB01")]
        assert _row_verdict(rows_list[0]) is True
        assert _row_verdict(rows_list[1]) is False
        assert _group_results(rows_list) == []

    def test_sin_par_valido_none_en_ambos_modos(self):
        rows_list = [_row_7h("890701")]
        del rows_list[0]["fecha_cierre"]
        assert _row_verdict(rows_list[0]) is False
        assert _group_results(rows_list) == []

    def test_puente_no_requiere_agg_en_param(self):
        """Sin compute_horas ni collect_set en agg_configs: el puente
        resuelve igual (cero migración de condiciones Y de params)."""
        rows_list = [_row_7h("890701")]
        results = _group_results(rows_list, agg_configs=[])
        assert len(results) == 1

    def test_convive_con_agg_existente(self):
        rows_list = [_row_7h("890701")]
        aggs = [{"function": "compute_horas", "field1": "fec_factura",
                 "field2": "fecha_cierre", "target": "estancia_horas"}]
        results = _group_results(rows_list, agg_configs=aggs)
        assert len(results) == 1

    def test_path_worksheet_tambien_puentea(self):
        wb = Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="NUMERO_FACTURA")
        ws.cell(row=1, column=2, value="CODIGO")
        ws.cell(row=1, column=3, value="TIPO_FACTURA_DESCRIPCION")
        ws.cell(row=1, column=4, value="FEC_FACTURA")
        ws.cell(row=1, column=5, value="FECHA_CIERRE")
        ws.cell(row=2, column=1, value="F001")
        ws.cell(row=2, column=2, value="890701")
        ws.cell(row=2, column=3, value="Urgencias")
        ws.cell(row=2, column=4, value="2026-06-24 08:00:00")
        ws.cell(row=2, column=5, value="2026-06-24 15:00:00")
        indices = {"numero_factura": 0, "codigo": 1,
                   "tipo_factura_descripcion": 2,
                   "fec_factura": 3, "fecha_cierre": 4}
        from app.services.engine.condition_evaluator import ConditionEvaluator
        from app.services.engine.evidence_collector import EvidenceCollector
        from app.services.engine.group_evaluator import GroupEvaluator

        ce = ConditionEvaluator()
        tree = ce.build_tree(_conds_44())
        groups = GroupEvaluator.build_groups(ws, indices)
        rule_info = {"id": 44, "version": 1, "dominio": "urgencias",
                     "nombre": "sala_obs_mayor_6_horas",
                     "descripcion": "Sala obs > 6h", "severidad": "error"}
        results = GroupEvaluator.evaluate(
            groups, ws, indices, [], tree, ce, rule_info,
            EvidenceCollector(), record_evidence=False,
        )
        assert len(results) == 1
        assert results[0]["factura"] == "F001"


class TestPuenteUnitario:
    def test_in_acepta_lista_any_match(self):
        from app.services.engine.evaluators import InEvaluator

        ev = InEvaluator()
        assert ev.evaluate({}, ["890701", "X"], SALA_CODES) is False
        assert ev.evaluate({}, ["890701", "5DSB01"], SALA_CODES) is True
        assert ev.evaluate({}, [], SALA_CODES) is False
        # path escalar intacto
        assert ev.evaluate({}, "5DSB01", SALA_CODES) is True
        assert ev.evaluate({}, "890701", SALA_CODES) is False

    def test_date_horas_grupo_primer_par_valido(self):
        from app.services.engine.context import EvaluationContext
        from app.services.engine.providers import DateProvider

        provider = DateProvider()
        ctx = EvaluationContext(
            invoice_data={},
            group_rows=[{"fec_factura": None, "fecha_cierre": None},
                        {"fec_factura": "2026-06-24 08:00:00",
                         "fecha_cierre": "2026-06-24 15:00:00"}],
        )
        assert provider.resolve("date.horas", ctx) == 7

    def test_date_horas_grupo_sin_par_valido_none(self):
        from app.services.engine.context import EvaluationContext
        from app.services.engine.providers import DateProvider

        provider = DateProvider()
        ctx = EvaluationContext(
            invoice_data={"fec_factura": "2026-06-24 08:00:00"},
            group_rows=[{"fec_factura": "2026-06-24 08:00:00"}],
        )
        assert provider.resolve("date.horas", ctx) is None
