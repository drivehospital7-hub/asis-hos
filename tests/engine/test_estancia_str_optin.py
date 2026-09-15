"""estancia_str opt-in: solo se genera cuando la regla lo pide.

Pedido usuario: "Si pongo estancia_str se hace la conversion, sino no."
Hoy _enrich_estancia_str agrega el texto a TODO item con estancia_horas.

Opt-in = detalle_a_campo == 'estancia_str'
      o detalle_b_campo == 'estancia_str'
      o descripcion_template contiene '{estancia_str}'.
Legacy con estancia_str de origen se respeta (no se borra).
Formatters sin la clave caen a su fallback (codigo), sin formatear
por su cuenta desde estancia_horas.
"""

from __future__ import annotations

from datetime import datetime, timedelta
from unittest.mock import MagicMock


def _ce():
    from app.services.engine.condition_evaluator import ConditionEvaluator
    return ConditionEvaluator()


def _collector():
    from app.services.engine.evidence_collector import EvidenceCollector
    return EvidenceCollector()


def _rule_info_sin_mapeo(**over):
    base = {
        "id": 61, "version": 1, "dominio": "urgencias",
        "nombre": "sin_mapeo", "descripcion": "Sin mapeo",
        "severidad": "error",
    }
    base.update(over)
    return base


def _always_true_tree():
    return _ce().build_tree([
        {"id": 1, "padre_id": None, "tipo": "atomic", "operador": "gt",
         "fuente_datos": "group.estancia_horas", "valor_esperado": 1, "orden": 0},
    ])


def _mk_group_ws(fec0: datetime, fec1: datetime):
    from openpyxl import Workbook
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="NUMERO_FACTURA")
    ws.cell(row=1, column=2, value="FEC_FACTURA")
    ws.cell(row=1, column=3, value="FECHA_CIERRE")
    ws.cell(row=2, column=1, value="F001")
    ws.cell(row=2, column=2, value=fec0)
    ws.cell(row=2, column=3, value=fec1)
    indices = {"numero_factura": 0, "fec_factura": 1, "fecha_cierre": 2}
    return ws, indices


def _agg_horas():
    return [
        {"function": "compute_horas", "field1": "fec_factura",
         "field2": "fecha_cierre", "target": "estancia_horas"},
    ]


def _eval_grupo(rule_info, rows=None):
    from app.services.engine.group_evaluator import GroupEvaluator
    base = datetime(2026, 1, 1, 8, 0, 0)
    ws, indices = _mk_group_ws(base, base + timedelta(hours=30, minutes=30))
    groups = GroupEvaluator.build_groups(ws, indices)
    kwargs = {}
    if rows is not None:
        kwargs["rows"] = rows
    return GroupEvaluator.evaluate(
        groups, ws, indices, _agg_horas(), _always_true_tree(),
        _ce(), rule_info, _collector(), record_evidence=False, **kwargs,
    )


class TestGrupoOptIn:
    def test_sin_mapeo_no_genera_str(self):
        results = _eval_grupo(_rule_info_sin_mapeo())
        assert len(results) == 1
        assert results[0].get("estancia_horas") == 30.5
        assert "estancia_str" not in results[0]

    def test_detalle_a_optin(self):
        results = _eval_grupo(_rule_info_sin_mapeo(detalle_a_campo="estancia_str"))
        assert len(results) == 1
        assert results[0].get("estancia_str") == "1d 6h"

    def test_template_optin(self):
        results = _eval_grupo(_rule_info_sin_mapeo(
            descripcion_template="Estancia {estancia_str} observada"))
        assert len(results) == 1
        assert results[0].get("estancia_str") == "1d 6h"

    def test_legacy_preservado_sin_mapeo(self):
        base = datetime(2026, 1, 1, 8, 0, 0)
        results = _eval_grupo(
            _rule_info_sin_mapeo(),
            rows=[{"numero_factura": "F001", "fec_factura": base,
                   "fecha_cierre": base + timedelta(hours=30, minutes=30),
                   "estancia_str": "custom"}],
        )
        assert len(results) == 1
        assert results[0].get("estancia_str") == "custom"


def _make_cond(**over):
    m = MagicMock()
    for k, v in over.items():
        setattr(m, k, v)
    return m


def _eval_row(rule, rows):
    from app.services.engine.engine import RuleEvaluationEngine
    root = {
        "id": 1, "regla_id": rule.id, "padre_id": None,
        "tipo": "atomic", "operador": "eq",
        "fuente_datos": "invoice.codigo",
        "valor_esperado": "5DSB01", "orden": 0,
    }
    session = MagicMock()
    q = MagicMock()
    q.filter.return_value = q
    q.order_by.return_value = q
    q.first.return_value = rule
    q.all.return_value = [_make_cond(**root)]
    session.query.return_value = q
    engine = RuleEvaluationEngine(session)
    return engine.evaluate_sheet(
        rule.nombre, None,
        {"numero_factura": 0, "codigo": 1},
        persist=False, rows=rows,
    )


def _regla_row(**over):
    from app.models import Regla
    rule = Regla(
        id=61, nombre="sin_mapeo", dominio="urgencias",
        estado="active", version=1, prioridad=10, severidad="error",
    )
    for k, v in over.items():
        setattr(rule, k, v)
    return rule


class TestRowOptIn:
    def test_sin_mapeo_no_genera_str(self):
        results = _eval_row(_regla_row(), [
            {"numero_factura": "F001", "codigo": "5DSB01", "estancia_horas": 30.5}])
        assert len(results) == 1
        assert results[0].get("estancia_horas") == 30.5
        assert "estancia_str" not in results[0]

    def test_template_optin(self):
        results = _eval_row(
            _regla_row(descripcion_template="Estancia {estancia_str}"),
            [{"numero_factura": "F001", "codigo": "5DSB01", "estancia_horas": 30.5}])
        assert len(results) == 1
        assert results[0].get("estancia_str") == "1d 6h"

    def test_legacy_preservado_sin_mapeo(self):
        results = _eval_row(_regla_row(), [
            {"numero_factura": "F001", "codigo": "5DSB01",
             "estancia_horas": 30.5, "estancia_str": "custom"}])
        assert len(results) == 1
        assert results[0].get("estancia_str") == "custom"


class TestFormattersFallback:
    def test_grupo_formatter_sin_clave_cae_a_codigo(self):
        from app.services.normalized_rows import _format_cups_equivalentes
        out = _format_cups_equivalentes(
            {"codigo": "C001", "procedimiento": "EQUIV",
             "estancia_horas": 30.5, "problema": "P"}, {})
        assert out["detalle"] == "C001"
        assert "Estancia:" not in out["detalle"]

    def test_legacy_block_sin_clave_cae_a_codigo(self):
        from app.services.normalized_rows import build_normalized_rows
        rows = build_normalized_rows(use_grupo_mapping=False, error_groups={
            "Cups Equivalentes": [{
                "factura": "FAC-001", "codigo": "C001", "procedimiento": "EQUIV",
                "estancia_horas": 30.5, "problema": "P",
            }]
        }, responsables_map={})
        assert rows[0]["detalle"] == "C001"
        assert "Estancia:" not in rows[0]["detalle"]
