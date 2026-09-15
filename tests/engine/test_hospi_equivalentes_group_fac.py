"""Tests for rule #69 hospi_equivalentes_group_fac (fixed: error-when-missing).

Rule #69 was INVERTED: it flagged ERROR when the invoice CONTAINED 38114 or
129B02. Correct semantics (user-confirmed): ERROR when the Hospitalización
invoice brings NEITHER code — at least one of the two is mandatory at
invoice level (OR, not ALL: set_intersects + NOT, hermanas pattern from
migration 010; the Hospitalización filter comes from the rule parametros).

Tests use direct condition trees (no DB) matching how
test_hospitalizacion_codes_rules.py works, plus unit tests for the
cups_equivalentes display helper in hospitalizacion/detect_all.py.
"""

from __future__ import annotations

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet


EQUIV_OBLIG = ["38114", "129B02"]


# ── Helpers ────────────────────────────────────────────────────────────────

def _build_ws() -> tuple[Worksheet, dict[str, int]]:
    """Create a worksheet with hospitalizacion columns."""
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="NUMERO_FACTURA")
    ws.cell(row=1, column=2, value="CODIGO")
    ws.cell(row=1, column=3, value="TIPO_FACTURA_DESCRIPCION")
    indices = {
        "numero_factura": 0,
        "codigo": 1,
        "tipo_factura_descripcion": 2,
    }
    return ws, indices


def _ce():
    from app.services.engine.condition_evaluator import ConditionEvaluator
    return ConditionEvaluator()


def _collector():
    from app.services.engine.evidence_collector import EvidenceCollector
    return EvidenceCollector()


def _rule_info(name="hospi_equivalentes_group_fac"):
    return {
        "id": 69, "version": 1, "dominio": "hospitalizacion",
        "nombre": name, "descripcion": "Test rule", "severidad": "error",
    }


def _fixed_tree() -> dict:
    """Corrected #69 tree: NOT[set_intersects(collect_set, 38114/129B02)]."""
    flat = [
        {"id": 1, "padre_id": None, "tipo": "composite", "operador": "NOT", "orden": 0},
        {"id": 2, "padre_id": 1, "tipo": "atomic", "operador": "set_intersects",
         "fuente_datos": "invoice.collect_set_codigo",
         "valor_esperado": EQUIV_OBLIG, "orden": 0},
    ]
    return _ce().build_tree(flat)


def _old_inverted_tree() -> dict:
    """Pre-fix #69 shape: OR[in(invoice.codigo)] — MATCHes when present."""
    flat = [
        {"id": 1, "padre_id": None, "tipo": "composite", "operador": "OR", "orden": 0},
        {"id": 2, "padre_id": 1, "tipo": "atomic", "operador": "in",
         "fuente_datos": "invoice.codigo",
         "valor_esperado": EQUIV_OBLIG, "orden": 0},
    ]
    return _ce().build_tree(flat)


def _agg() -> list[dict]:
    return [
        {"function": "collect_set", "field": "codigo", "target": "collect_set_codigo"},
    ]


def _evaluate(ws, indices, tree) -> list[dict]:
    from app.services.engine.group_evaluator import GroupEvaluator
    groups = GroupEvaluator.build_groups(
        ws, indices, filter_field="tipo_factura_descripcion",
        filter_value="Hospitalización",
    )
    return GroupEvaluator.evaluate(
        groups, ws, indices, _agg(), tree, _ce(),
        _rule_info(), _collector(),
    )


# ═══════════════════════════════════════════════════════════════════════════
# Fixed tree: error-when-missing (OR semantics)
# ═══════════════════════════════════════════════════════════════════════════

class TestEquivGroupFacFixed:
    """Corrected #69: ERROR only when NEITHER 38114 nor 129B02 is present."""

    def test_129b02_en_segunda_fila_no_error(self):
        """129B02 in the 2nd row of the factura → NO error (was false +)."""
        ws, indices = _build_ws()
        ws.cell(row=2, column=1, value="F001")
        ws.cell(row=2, column=2, value="999999")
        ws.cell(row=2, column=3, value="Hospitalización")
        ws.cell(row=3, column=1, value="F001")
        ws.cell(row=3, column=2, value="129B02")
        ws.cell(row=3, column=3, value="Hospitalización")

        assert _evaluate(ws, indices, _fixed_tree()) == []

    def test_38114_presente_no_error(self):
        """38114 present anywhere in the factura → NO error."""
        ws, indices = _build_ws()
        ws.cell(row=2, column=1, value="F001")
        ws.cell(row=2, column=2, value="38114")
        ws.cell(row=2, column=3, value="Hospitalización")

        assert _evaluate(ws, indices, _fixed_tree()) == []

    def test_sin_ninguno_error(self):
        """Factura with neither code → exactly 1 error for F001."""
        ws, indices = _build_ws()
        ws.cell(row=2, column=1, value="F001")
        ws.cell(row=2, column=2, value="999999")
        ws.cell(row=2, column=3, value="Hospitalización")
        ws.cell(row=3, column=1, value="F001")
        ws.cell(row=3, column=2, value="888888")
        ws.cell(row=3, column=3, value="Hospitalización")

        results = _evaluate(ws, indices, _fixed_tree())
        assert len(results) == 1
        assert results[0]["factura"] == "F001"

    def test_urgencias_sin_ninguno_filtrada(self):
        """Urgencias factura missing both codes → filtered, NO error."""
        ws, indices = _build_ws()
        ws.cell(row=2, column=1, value="F001")
        ws.cell(row=2, column=2, value="999999")
        ws.cell(row=2, column=3, value="Urgencias")

        assert _evaluate(ws, indices, _fixed_tree()) == []

    def test_old_tree_was_inverted(self):
        """Pre-fix shape MATCHes when the code IS present (documents the bug)."""
        ws, indices = _build_ws()
        ws.cell(row=2, column=1, value="F001")
        ws.cell(row=2, column=2, value="129B02")
        ws.cell(row=2, column=3, value="Hospitalización")

        results = _evaluate(ws, indices, _old_inverted_tree())
        assert len(results) == 1
        assert results[0]["factura"] == "F001"


# ═══════════════════════════════════════════════════════════════════════════
# Display helper: intersection, never the first-row codigo
# ═══════════════════════════════════════════════════════════════════════════

class TestEquivDisplay:
    """_codigo_equivalentes_display shows the 38114/129B02 intersection."""

    def test_muestra_interseccion_no_primera_fila(self):
        """collect_set with 38114 + first-row codigo noise → '38114'."""
        from app.services.hospitalizacion.detect_all import (
            _codigo_equivalentes_display,
        )
        item = {
            "factura": "F001",
            "codigo": "999999",  # first-row noise from the parity bridge
            "collect_set_codigo": ["999999", "38114"],
        }
        assert _codigo_equivalentes_display(item) == "38114"

    def test_sin_interseccion_muestra_collect_set(self):
        """Missing-code error → shows what the factura actually brought."""
        from app.services.hospitalizacion.detect_all import (
            _codigo_equivalentes_display,
        )
        item = {
            "factura": "F001",
            "codigo": "999999",
            "collect_set_codigo": ["999999", "888888"],
        }
        assert _codigo_equivalentes_display(item) == "888888, 999999"

    def test_legacy_list_codigo_fallback(self):
        """Legacy row-shape items with list codigo still render."""
        from app.services.hospitalizacion.detect_all import (
            _codigo_equivalentes_display,
        )
        item = {"factura": "F001", "codigo": ["890601H", "129B02"]}
        assert _codigo_equivalentes_display(item) == "129B02, 890601H"
