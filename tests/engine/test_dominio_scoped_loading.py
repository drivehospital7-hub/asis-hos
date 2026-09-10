"""TDD tests for dominio-grupo-error T4: dominio-scoped rule loading.

Spec domain-scoped-evaluation:
- exact domain hit wins (highest version, requested dominio)
- duplicate (nombre, dominio) picks highest version deterministically
- NULL-dominio row never matches (no fallback)
- transversal rules resolve from any area
"""

from __future__ import annotations

from unittest.mock import MagicMock


def _make_rule(rule_id, nombre, dominio, version=1):
    from app.models import Regla

    return Regla(
        id=rule_id,
        nombre=nombre,
        dominio=dominio,
        estado="active",
        version=version,
        prioridad=100,
        severidad="error",
        activo=True,
    )


def _session_returning(rule_or_none):
    session = MagicMock()
    query = MagicMock()
    query.filter.return_value = query
    query.order_by.return_value = query
    query.first.return_value = rule_or_none
    session.query.return_value = query
    return session


class TestLoadRuleByNameDominio:
    def test_exact_domain_hit_highest_version(self):
        from app.services.engine.engine import RuleEvaluationEngine

        winner = _make_rule(2, "codigo_entidad", "urgencias", version=3)
        engine = RuleEvaluationEngine(_session_returning(winner))
        loaded = engine._load_rule_by_name("codigo_entidad", "urgencias")
        assert loaded is not None
        assert loaded.version == 3
        assert loaded.dominio == "urgencias"

    def test_filter_uses_dominio_and_orders_version_desc(self):
        from app.services.engine.engine import RuleEvaluationEngine

        winner = _make_rule(5, "x_rule", "odontologia", version=2)
        session = _session_returning(winner)
        engine = RuleEvaluationEngine(session)
        engine._load_rule_by_name("x_rule", "odontologia")
        query = session.query.return_value
        assert query.order_by.called, "must order by version DESC deterministically"

    def test_null_dominio_row_never_matches(self):
        from app.services.engine.engine import RuleEvaluationEngine

        engine = RuleEvaluationEngine(_session_returning(None))
        assert engine._load_rule_by_name("x_rule", "urgencias") is None

    def test_transversal_resolves_from_any_area(self):
        from app.services.engine.engine import RuleEvaluationEngine

        transversal = _make_rule(9, "valores_decimales", "transversal", version=1)
        engine = RuleEvaluationEngine(_session_returning(transversal))
        loaded = engine._load_rule_by_name("valores_decimales", "farmacia")
        assert loaded is not None
        assert loaded.dominio == "transversal"


class TestDuplicateVersionRealDb:
    """Real-SQLite triangulation: ordering/exact-priority executes for real."""

    def _engine_with_rows(self, rowspec):
        from sqlalchemy import create_engine
        from sqlalchemy.orm import sessionmaker
        from sqlalchemy.pool import StaticPool
        from app.database import Base
        from app.models import Regla
        import app.models  # noqa: F401
        from app.services.engine.engine import RuleEvaluationEngine

        engine = create_engine(
            "sqlite://",
            connect_args={"check_same_thread": False},
            poolclass=StaticPool,
        )
        Base.metadata.create_all(engine)
        session = sessionmaker(bind=engine)()
        for i, (nombre, dominio, version) in enumerate(rowspec):
            session.add(
                Regla(
                    nombre=nombre, dominio=dominio, estado="active",
                    version=version, prioridad=100, severidad="error",
                    activo=True,
                )
            )
        session.commit()
        return RuleEvaluationEngine(session)

    def test_same_nombre_dominio_picks_highest_version(self):
        engine = self._engine_with_rows([
            ("codigo_entidad", "urgencias", 1),
            ("codigo_entidad", "urgencias", 2),
        ])
        loaded = engine._load_rule_by_name("codigo_entidad", "urgencias")
        assert loaded.version == 2

    def test_exact_dominio_beats_higher_transversal_version(self):
        engine = self._engine_with_rows([
            ("codigo_entidad", "odontologia", 2),
            ("codigo_entidad", "transversal", 9),
        ])
        loaded = engine._load_rule_by_name("codigo_entidad", "odontologia")
        assert loaded.dominio == "odontologia"
        assert loaded.version == 2

    def test_transversal_used_when_no_exact_match(self):
        engine = self._engine_with_rows([
            ("codigo_entidad", "odontologia", 5),
            ("codigo_entidad", "transversal", 3),
        ])
        loaded = engine._load_rule_by_name("codigo_entidad", "farmacia")
        assert loaded.dominio == "transversal"


class TestEvaluateSheetThreadsDominio:
    def test_evaluate_sheet_accepts_dominio_kwarg(self):
        from app.services.engine.engine import RuleEvaluationEngine
        from openpyxl import Workbook

        winner = _make_rule(1, "valores_decimales", "transversal", version=1)
        session = _session_returning(winner)
        # No conditions -> empty tree -> returns []
        query_all = session.query.return_value
        query_all.all.return_value = []

        wb = Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="NUMERO_FACTURA")
        ws.cell(row=2, column=1, value="F001")

        engine = RuleEvaluationEngine(session)
        results = engine.evaluate_sheet(
            "valores_decimales", ws, {"numero_factura": 0}, dominio="urgencias"
        )
        assert results == []
