"""Migration 015 tests (strict TDD, SQL-text + hermetic firing — zero DB writes).

015 seeds 4 of the 7 intramural GAP rules (Ref #1) by (nombre, version),
idempotent ON CONFLICT + DELETE+rebuild like 010/011/012/013/014. Never
hardcodes live row IDs, never BEGIN/COMMIT or MAX(id).

Seeded (4, all dominio='intramural', active v1):
  bacteriologas_cronograma (1 cond cronograma_check),
  centro_costo_intramural_valido (100 conds = F15 OR tree verbatim),
  duplicado_id_codigo_05 (1 cond gte count>=2 + group_by params),
  revision_cantidad_intramural (1 cond cascade evaluator).
Total: 103 condiciones.

Documented skips (3 — need product input, must stay absent here):
  ide_contrato_simple (catalog 'ide_simple_rules' seeded nowhere),
  pym_rutas_dx (pre_scan_sheet never called by the engine → dead rule),
  duplicado_id_codigo_02_lab (needs tipo=02 AND lab=Si; group params allow
  a single filter_field only).

Firing tests serve the 015-seeded shapes via mock sessions (same _run pattern
as test_dangling_rule_refs) on the code's own fixture shapes. Hermetic:
MagicMock sessions + mocked cronograma service — no DB writes, no servers.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any
from unittest.mock import MagicMock, patch

from openpyxl import Workbook

MIGRATIONS_DIR = Path("migrations")
MIGRATION = MIGRATIONS_DIR / "015_seed_intramural_gaps.sql"
SEED_F15 = Path("seed/migracion-engine/15_centro_costo_intramural.sql")

EXPECTED_RULES = {
    "bacteriologas_cronograma": "intramural",
    "centro_costo_intramural_valido": "intramural",
    "duplicado_id_codigo_05": "intramural",
    "revision_cantidad_intramural": "intramural",
}

EXPECTED_COUNTS = {
    "bacteriologas_cronograma": 1,
    "centro_costo_intramural_valido": 100,
    "duplicado_id_codigo_05": 1,
    "revision_cantidad_intramural": 1,
}

# The 3 documented skips — product input required, must NOT be seeded here.
DOCUMENTED_SKIPS = (
    "ide_contrato_simple",
    "pym_rutas_dx",
    "duplicado_id_codigo_02_lab",
)

# All 12 cat_in keys used by the F15 tree — seeded by 011/012/013, never here.
F15_CATALOG_KEYS = (
    "codigos_exceptuados",
    "centro_costo_quirofano",
    "centro_costo_hospitalizacion",
    "centro_costo_pyp",
    "centros_costo_pyp_intramural",
    "codigos_tipo_procedimiento_laboratorio",
    "centros_costo_laboratorio_validos",
    "codigos_excluidos_vacunacion",
    "codigos_tipo_procedimiento_ambulatorio",
    "codigos_exceptuados_ambulatorio",
    "facturadores_urgencias",
    "codigos_exceptuados_responsable_urgencias",
)


def _text() -> str:
    return MIGRATION.read_text(encoding="utf-8")


def _code_lines(text: str) -> list[str]:
    """SQL lines with full-line comments stripped (audits ignore comments)."""
    return [ln for ln in text.splitlines() if not ln.strip().startswith("--")]


def _rule_window(text: str, name: str) -> str:
    """Executable window for one rule: from its upsert to the next rule block."""
    start = text.find(f"'{name}', '")
    assert start != -1, f"rule {name} upsert missing"
    rest = text[start + len(name):]
    nxt = re.search(r"\n-- \d\. ", rest)
    return rest[: nxt.start()] if nxt else rest


# ── SQL-text tests (same contract as 011/012/013/014) ──


def test_015_file_exists() -> None:
    assert MIGRATION.exists(), "015 migration file missing"


def test_015_planned_after_014() -> None:
    from run_migrations import plan_migrations

    planned = [p.stem for p in plan_migrations(MIGRATIONS_DIR, applied=set(), include_evidence=False)]
    assert "014_seed_final_unseeded_rules" in planned
    assert "015_seed_intramural_gaps" in planned
    assert planned.index("014_seed_final_unseeded_rules") < planned.index(
        "015_seed_intramural_gaps"
    )


def test_015_not_treated_as_evidence_seed() -> None:
    from run_migrations import should_skip_evidence

    assert should_skip_evidence(MIGRATION.name, include_evidence=False) is False
    assert should_skip_evidence(MIGRATION.name, include_evidence=True) is False


def test_015_upserts_each_rule_by_name_version() -> None:
    text = _text()
    for name, dominio in EXPECTED_RULES.items():
        pattern = re.compile(
            r"'%s',\s*'.*?',\s*'%s',\s*'active',\s*1," % (re.escape(name), re.escape(dominio)),
            re.DOTALL,
        )
        assert pattern.search(text), f"rule {name} not upserted as ('name', {dominio}, active, v1)"
    assert text.count("ON CONFLICT (nombre, version) DO UPDATE") >= len(EXPECTED_RULES)


def test_015_rebuilds_conditions_by_name() -> None:
    text = _text()
    for name in EXPECTED_RULES:
        assert re.search(
            r"WHERE nombre = '%s' AND version = 1" % re.escape(name), text
        ), f"rule {name} conditions not resolved by (nombre, version)"
    assert text.count("DELETE FROM condiciones WHERE regla_id = ") >= len(EXPECTED_RULES)


def test_015_no_hardcoded_rule_ids() -> None:
    code = "\n".join(_code_lines(_text()))
    assert re.search(r"regla_id\s*=\s*\d+", code) is None, "hardcoded regla_id literal found"
    assert re.search(r"padre_id\s*=\s*\d+", code) is None, "hardcoded padre_id literal found"


def test_015_no_transaction_control_or_max_id() -> None:
    code = "\n".join(_code_lines(_text()))
    assert re.search(r"^\s*BEGIN\s*;", code, re.MULTILINE) is None
    assert re.search(r"^\s*COMMIT\s*;", code, re.MULTILINE) is None
    assert "MAX(id)" not in code and "MAX(c.id)" not in code


def test_015_seeds_expected_cond_counts() -> None:
    """Per-rule INSERT counts: 1 + 100 + 1 + 1 = 103."""
    text = _text()
    for name, expected in EXPECTED_COUNTS.items():
        window = _rule_window(text, name)
        found = window.count("VALUES (_regla_id,")
        assert found == expected, f"rule {name}: seeded {found} conds, expected {expected}"
    assert text.count("VALUES (_regla_id,") == sum(EXPECTED_COUNTS.values()) == 103


def test_015_documented_skips_absent() -> None:
    """The 3 ambiguous rules need product input — none may be seeded here."""
    code = "\n".join(_code_lines(_text()))
    for name in DOCUMENTED_SKIPS:
        assert f"'{name}'" not in code, f"ambiguous rule {name} must stay unseeded (needs product input)"


def test_015_seeds_no_catalogs() -> None:
    """Zero catalog writes: the F15 tree's 12 cat_in keys come from 011/012/013."""
    code = "\n".join(_code_lines(_text()))
    assert "INSERT INTO catalogos" not in code


def test_015_single_active_v1_per_rule() -> None:
    text = _text()
    assert "version = 2" not in text
    assert "version <> 1" not in text
    for name in EXPECTED_RULES:
        assert text.count(f"WHERE nombre = '{name}' AND version = 1") == 1


def test_015_guarded_schema_alters_like_010() -> None:
    lowered = _text().lower()
    assert "information_schema" in lowered
    assert "operador" in lowered and "valor_esperado" in lowered


def test_015_backfills_rule_base_id() -> None:
    text = _text()
    assert "rule_base_id" in text
    assert re.search(r"rule_base_id\s*=\s*id", text) is not None
    for name in EXPECTED_RULES:
        assert name in text


# ── F15-alignment tests ──


def test_015_centro_tree_matches_f15_verbatim() -> None:
    """The centro window carries the F15 OR tree unchanged (F15 aligned)."""
    text = _text()
    window = _rule_window(text, "centro_costo_intramural_valido")
    f15 = SEED_F15.read_text(encoding="utf-8")
    f15_code = "\n".join(ln for ln in f15.splitlines() if not ln.strip().startswith("--"))

    def norm(s: str) -> str:
        return re.sub(r"\s+", " ", s).strip()

    w_inserts = [norm(ln) for ln in window.splitlines() if "VALUES (_regla_id," in ln]
    f_inserts = [norm(ln) for ln in f15_code.splitlines() if "VALUES (_regla_id," in ln]
    assert w_inserts == f_inserts, "015 centro tree drifted from seed/migracion-engine/15"


def test_015_centro_tree_shape() -> None:
    """OR root + 18 AND branches + nested composites; all catalog keys used."""
    text = _text()
    window = _rule_window(text, "centro_costo_intramural_valido")
    assert window.count("'OR'") == 3, "root OR + 2 nested OR (rev10, rev6)"
    assert window.count("'AND'") == 21, "18 root branches + 3 nested AND"
    assert window.count("'NOT'") == 26
    assert window.count("'cat_in'") == 22
    for key in F15_CATALOG_KEYS:
        assert f"'{key}'" in window, f"F15 tree missing catalog key {key}"
    # Intramural-specific markers (literal center/IDE values from the tree)    assert '"URGENCIAS","HOSPITALIZACIÓN - ESTANCIA GENERAL"' in window.replace(" ", "")
    assert "SALUD PUBLICA-VACUNACION  REGULAR" in window
    assert "SERVICIOS AMBULATORIOS- CONSULTA EXTERNA Y PROCEDIMIENTOS" in window


def test_015_single_cond_shapes() -> None:
    """The 3 single-cond rules use their code-registered operators/fuentes."""
    text = _text()
    bact = _rule_window(text, "bacteriologas_cronograma")
    assert "'cronograma_check', 'invoice.codigo_profesional'" in bact
    dup = _rule_window(text, "duplicado_id_codigo_05")
    assert "'gte', 'invoice.count', '2'" in dup
    assert '"group_by": ["identificacion", "codigo", "codigo_dx_principal"]' in dup
    assert '"filter_field": "codigo_tipo_procedimiento", "filter_value": "05"' in dup
    assert '"function": "group_size", "target": "count"' in dup
    assert '"function": "collect_group_keys", "field": "numero_factura"' in dup
    rev = _rule_window(text, "revision_cantidad_intramural")
    assert "'revision_cantidad_intramural', 'invoice.cantidad'" in rev


# ── Hermetic firing tests (mock session serves the 015 shape) ──


def _mock_session(
    rule_name: str | None,
    dominio: str,
    descripcion: str,
    condiciones: list[dict],
    parametros: Any = None,
):
    """Session serving one rule tree, or nothing when rule_name is None."""
    from app.models import Regla

    session = MagicMock()
    mock_query = MagicMock()
    mock_query.filter.return_value = mock_query
    mock_query.order_by.return_value = mock_query

    if rule_name is None:
        mock_query.first.return_value = None
        mock_query.all.return_value = []
    else:
        regla = Regla(
            id=1, nombre=rule_name, dominio=dominio,
            estado="active", version=1, prioridad=10, severidad="warning",
            descripcion=descripcion,
        )
        regla.parametros = parametros
        mock_query.first.return_value = regla
        cond_mocks = []
        for cd in condiciones:
            m = MagicMock()
            m.id = cd["id"]
            m.regla_id = cd.get("regla_id", 1)
            m.padre_id = cd["padre_id"]
            m.tipo = cd["tipo"]
            m.operador = cd.get("operador")
            m.fuente_datos = cd.get("fuente_datos")
            m.valor_esperado = cd.get("valor_esperado")
            m.orden = cd.get("orden", 0)
            cond_mocks.append(m)
        mock_query.all.return_value = cond_mocks
    session.query.return_value = mock_query
    return session


def _run(rule_name: str | None, dominio: str, descripcion: str,
         condiciones: list[dict], ws, indices, parametros: Any = None):
    from app.services.engine.rule_based_detector import RuleBasedDetector

    session = _mock_session(rule_name, dominio, descripcion, condiciones, parametros)
    return RuleBasedDetector(rule_name or "dangling", session, dominio=dominio).detect(
        ws, indices, persist=False,
    )


def _wb(headers: list[str], rows: list[tuple]) -> Workbook:
    wb = Workbook()
    ws = wb.active
    for col, name in enumerate(headers, start=1):
        ws.cell(row=1, column=col, value=name)
    for r, row in enumerate(rows, start=2):
        for c, val in enumerate(row, start=1):
            ws.cell(row=r, column=c, value=val)
    return wb


def _cid_seq():
    n = [0]

    def nxt() -> int:
        n[0] += 1
        return n[0]

    return nxt


class TestBacteriologasFires:
    """015 cronograma_check tree fires on the evaluator's own fixture shape
    (03374 = BACTERIOLOGA, tipo 02 + Lab Si, not excepted, empty turnos)."""

    HEADERS = ["numero_factura", "codigo_profesional", "tipo_factura_descripcion",
               "codigo_tipo_procedimiento", "laboratorio", "codigo",
               "responsable_cierra", "fec_factura"]

    def _tree(self):
        return [{"id": 1, "padre_id": None, "tipo": "atomic",
                 "operador": "cronograma_check",
                 "fuente_datos": "invoice.codigo_profesional",
                 "valor_esperado": None, "orden": 0}]

    def _fixture(self):
        wb = _wb(self.HEADERS, [("I-001", "03374", "Intramural", "02", "Si",
                                 "901210", "", "2024-06-01")])
        return wb, {h: i for i, h in enumerate(self.HEADERS)}

    def test_fires_when_off_turno(self):
        from app.services.engine.evaluators import EVALUATOR_REGISTRY

        EVALUATOR_REGISTRY["cronograma_check"]._cronograma_cache.clear()
        wb, indices = self._fixture()
        with patch(
            "app.services.cronograma_bacteriologas_service.get_turno_del_dia",
            return_value=[{"nombre": "OTRA PERSONA", "codigo": "CE"}],
        ):
            res = _run("bacteriologas_cronograma", "intramural",
                       "cronograma del dia", self._tree(), wb.active, indices)
        assert len(res) == 1 and res[0]["factura"] == "I-001"

    def test_quiet_when_on_turno(self):
        from app.services.engine.evaluators import EVALUATOR_REGISTRY

        EVALUATOR_REGISTRY["cronograma_check"]._cronograma_cache.clear()
        wb, indices = self._fixture()
        with patch(
            "app.services.cronograma_bacteriologas_service.get_turno_del_dia",
            return_value=[{"nombre": "MOLINA ALVAREZ KAROL DAYANNA",
                           "codigo": "CE/PYM"}],
        ):
            res = _run("bacteriologas_cronograma", "intramural",
                       "cronograma del dia", self._tree(), wb.active, indices)
        assert res == []


class TestCentroRegla9BranchFires:
    """REGLA9 branch verbatim from 015 fires: tarifario farmacia fuera de
    farmacia (same fixture as test_intramural_engine_f4 T-F4.1)."""

    HEADERS = ["numero_factura", "centro_costo", "codigo",
               "codigo_tipo_procedimiento", "laboratorio", "tarifario",
               "responsable_cierra"]

    def _tree(self):
        nid = _cid_seq()
        root = {"id": nid(), "padre_id": None, "tipo": "composite",
                "operador": "AND", "orden": 0}
        notn = {"id": nid(), "padre_id": root["id"], "tipo": "composite",
                "operador": "NOT", "orden": 1}
        return [root,
                {"id": nid(), "padre_id": root["id"], "tipo": "atomic",
                 "operador": "eq", "fuente_datos": "invoice.tarifario",
                 "valor_esperado": "Suminstros, Medicamentos", "orden": 0},
                notn,
                {"id": nid(), "padre_id": notn["id"], "tipo": "atomic",
                 "operador": "eq", "fuente_datos": "invoice.centro_costo",
                 "valor_esperado": "APOYO TERAPEUTICO-FARMACIA E INSUMOS.",
                 "orden": 0}]

    def test_fires_farmacia_tarifario_fuera_farmacia(self):
        wb = _wb(self.HEADERS, [("I-002", "URGENCIAS", "890201", "01", "No",
                                 "Suminstros, Medicamentos", "")])
        indices = {h: i for i, h in enumerate(self.HEADERS)}
        res = _run("centro_costo_intramural_valido", "intramural",
                   "centro costo intramural", self._tree(), wb.active, indices)
        assert len(res) == 1 and res[0]["factura"] == "I-002"

    def test_quiet_farmacia_tarifario_en_farmacia(self):
        wb = _wb(self.HEADERS, [("I-002", "APOYO TERAPEUTICO-FARMACIA E INSUMOS.",
                                 "890201", "01", "No",
                                 "Suminstros, Medicamentos", "")])
        indices = {h: i for i, h in enumerate(self.HEADERS)}
        res = _run("centro_costo_intramural_valido", "intramural",
                   "centro costo intramural", self._tree(), wb.active, indices)
        assert res == []


class TestDuplicado05Fires:
    """015 dup_05 shape (seed 07 params + gte count 2) fires when one
    (ident, codigo, dx) group of tipo 05 appears twice."""

    HEADERS = ["numero_factura", "identificacion", "codigo",
               "codigo_dx_principal", "codigo_tipo_procedimiento"]

    PARAMS = [{"group_by": ["identificacion", "codigo", "codigo_dx_principal"],
               "filter_field": "codigo_tipo_procedimiento", "filter_value": "05",
               "aggregations": [{"function": "group_size", "target": "count"},
                                {"function": "collect_group_keys",
                                 "field": "numero_factura", "target": "facturas"}]}]

    def _tree(self):
        return [{"id": 1, "padre_id": None, "tipo": "atomic", "operador": "gte",
                 "fuente_datos": "invoice.count", "valor_esperado": "2",
                 "orden": 0}]

    def test_fires_on_duplicate_group(self):
        wb = _wb(self.HEADERS, [("I-101", "111", "A123", "J00", "05"),
                                ("I-102", "111", "A123", "J00", "05")])
        indices = {h: i for i, h in enumerate(self.HEADERS)}
        res = _run("duplicado_id_codigo_05", "intramural", "duplicados 05",
                   self._tree(), wb.active, indices, parametros=self.PARAMS)
        assert len(res) == 1

    def test_quiet_on_single_row(self):
        wb = _wb(self.HEADERS, [("I-101", "111", "A123", "J00", "05")])
        indices = {h: i for i, h in enumerate(self.HEADERS)}
        res = _run("duplicado_id_codigo_05", "intramural", "duplicados 05",
                   self._tree(), wb.active, indices, parametros=self.PARAMS)
        assert res == []


class TestRevisionCantidadFires:
    """015 revision cascade fires on the legacy boundary: tipo 02 + Lab No
    flags cantidad 3 (> CANTIDAD_MAX_02_NO_LAB=2), stays quiet at 2."""

    HEADERS = ["numero_factura", "cantidad", "codigo",
               "codigo_tipo_procedimiento", "laboratorio"]

    def _tree(self):
        return [{"id": 1, "padre_id": None, "tipo": "atomic",
                 "operador": "revision_cantidad_intramural",
                 "fuente_datos": "invoice.cantidad",
                 "valor_esperado": None, "orden": 0}]

    def _run_cantidad(self, cantidad: int):
        wb = _wb(self.HEADERS, [("I-003", cantidad, "901210", "02", "No")])
        indices = {h: i for i, h in enumerate(self.HEADERS)}
        return _run("revision_cantidad_intramural", "intramural",
                    "revision cantidad", self._tree(), wb.active, indices)

    def test_fires_above_threshold(self):
        res = self._run_cantidad(3)
        assert len(res) == 1 and res[0]["factura"] == "I-003"

    def test_quiet_at_threshold(self):
        assert self._run_cantidad(2) == []
