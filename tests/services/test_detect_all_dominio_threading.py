"""TDD tests for dominio-grupo-error T6 + Cambio 2: dominio threading.

Spec area isolation parity: dynamic discovery (RuleResolver + loop
RuleBasedDetector in domain_detection) threads the AREA dominio into every
detector; no area orchestrator hardcodes rule names. An urgencias sheet with
an odontologia-only violation yields zero findings from that rule.
"""

from __future__ import annotations

import re
from pathlib import Path
from unittest.mock import MagicMock, patch

SERVICES_DIR = Path(__file__).resolve().parent.parent.parent / "app" / "services"

AREA_MODULES = {
    "ambulatoria": "ambulatoria",
    "equipos_basicos": "equipos_basicos",
    "extramural": "extramural",
    "farmacia": "farmacia",
    "hospitalizacion": "hospitalizacion",
    "intramural": "intramural",
    "odontologia": "odontologia",
    "urgencias": "urgencias",
}

HARDCODED_CALL = re.compile(r'RuleBasedDetector\(\s*"[^"]+"\s*,')
HARDCODED_LOOP = re.compile(r'RuleBasedDetector\(\s*rule_name\s*,')


def _read(area: str) -> str:
    return (SERVICES_DIR / area / "detect_all.py").read_text(encoding="utf-8")


class TestNoHardcodedRuleNames:
    def test_no_literal_detector_calls(self):
        for area in AREA_MODULES:
            literal = HARDCODED_CALL.findall(_read(area))
            assert not literal, f"{area}/detect_all.py hardcodes rule names: {literal[:3]}"

    def test_no_variable_detector_calls(self):
        for area in AREA_MODULES:
            looped = HARDCODED_LOOP.findall(_read(area))
            assert not looped, f"{area}/detect_all.py loops RuleBasedDetector by name: {looped[:3]}"

    def test_each_area_discovers_via_domain_detection(self):
        for area, dominio in AREA_MODULES.items():
            text = _read(area)
            const = f"AREA_{dominio.upper()}"
            assert "detect_domain_rules" in text, (
                f"{area}/detect_all.py does not use dynamic discovery"
            )
            assert const in text, (
                f"{area}/detect_all.py does not reference {const}"
            )


class TestAreaIsolationBehavioral:
    def _sheet(self):
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        for col, header in enumerate(
            ["NUMERO_FACTURA", "RESPONSABLE_CIERRA", "FECHA_CIERRE", "FEC_FACTURA"],
            start=1,
        ):
            ws.cell(row=1, column=col, value=header)
        ws.cell(row=2, column=1, value="F001")
        ws.cell(row=2, column=2, value="ALGUIEN")
        return ws

    def test_urgencias_detectors_all_receive_urgencias_dominio(self):
        """Dynamic discovery threads the AREA dominio into every detector."""
        from app.models import Regla

        from app.services.urgencias import detect_all as urgencias_mod

        captured: list[str] = []

        class _FakeDetector:
            def __init__(self, rule_name, session, dominio):
                captured.append(dominio)
                self._rule_name = rule_name

            def detect(self, *args, **kwargs):
                # Odontologia-only violation: no urgencias rule fires
                return []

        served = [
            Regla(
                id=1, nombre="valores_decimales", dominio="transversal",
                estado="active", version=1, prioridad=10, severidad="error",
                activo=True, grupo_error="Decimales",
            ),
            Regla(
                id=2, nombre="nueva_regla_ui", dominio="urgencias",
                estado="active", version=1, prioridad=20, severidad="error",
                activo=True, grupo_error="Decimales",
            ),
        ]
        fake_resolver = MagicMock()
        fake_resolver.resolve.return_value = served

        session = MagicMock()
        session_manager = MagicMock()
        session_manager.__enter__.return_value = session

        indices = {
            "numero_factura": 0,
            "responsable_cierra": 1,
            "fecha_cierre": 2,
            "fec_factura": 3,
        }
        with (
            patch(
                "app.services.engine.session_manager.SessionManager",
                return_value=session_manager,
            ),
            patch(
                "app.services.engine.domain_detection.RuleResolver",
                return_value=fake_resolver,
            ),
            patch(
                "app.services.engine.rule_based_detector.RuleBasedDetector",
                _FakeDetector,
            ),
        ):
            resultado, _ = urgencias_mod.detect_all_problems_urgencias(
                self._sheet(), indices
            )

        assert captured, "expected RuleBasedDetector instantiations"
        assert set(captured) == {"urgencias"}, f"leaked dominios: {set(captured)}"
        assert resultado["area"] == "urgencias"
