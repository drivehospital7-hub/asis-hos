"""Tests de ruta para app/routes/cruce_produccion.py.

El blueprint solo delega: se mockea ``detect_all`` (sin red) y
``get_roots``; el reader/matcher/exporter reales hacen el trabajo.
"""

from __future__ import annotations

import io

import pytest
from openpyxl import Workbook

import app.routes.cruce_produccion as routes
from app.services.monitoreo_carpetas import InvoiceRecord, ScanResult


def _login(app_client, permisos):
    with app_client.session_transaction() as sess:
        sess["ce_authenticated"] = True
        sess["username"] = "tester"
        sess["permisos"] = permisos


def _produccion_bytes() -> bytes:
    """Excel mínimo: 1 fila con carpeta existente + 1 sin carpeta."""
    wb = Workbook()
    ws = wb.active
    ws.cell(row=1, column=1, value="Número Factura")
    ws.cell(row=1, column=2, value="Responsable Cierra Facturar")
    ws.cell(row=1, column=3, value="Fec. Factura")
    ws.cell(row=2, column=1, value="FEV123")
    ws.cell(row=2, column=2, value="Ana")
    ws.cell(row=2, column=3, value="2026-01-05")
    ws.cell(row=3, column=1, value="FEV999")
    ws.cell(row=3, column=2, value="Luis")
    ws.cell(row=3, column=3, value="2026-02-01")
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def _scan_result() -> ScanResult:
    return ScanResult(
        facturas=[
            InvoiceRecord(
                filename="FEV123",
                facturador="Fact1",
                full_path="/root/Fact1/FEV123",
                status="En revisión",
                invoice_type="FEV",
                invoice_code="FEV123",
            )
        ],
        indicadores={},
        duplicados=[],
        vacias=[],
        errores_scan=[],
        excel_path=None,
    )


@pytest.fixture
def _mock_scan(monkeypatch):
    """Roots fijas + scan fresco simulado (sin tocar red) + sin novedades."""
    monkeypatch.setattr(routes, "get_roots", lambda: (["//fake-root"], "env", None))
    monkeypatch.setattr(routes, "detect_all", lambda _roots: _scan_result())
    monkeypatch.setattr(routes, "obtener_novedades", lambda: [])


# ── Shell ──


def test_shell_200_con_permiso(app_client):
    _login(app_client, ["cruce_produccion"])
    resp = app_client.get("/cruce-produccion/")
    assert resp.status_code == 200
    assert b"Cruce Producci" in resp.data


def test_shell_redirect_sin_permiso(app_client):
    _login(app_client, ["procesar"])
    resp = app_client.get("/cruce-produccion/")
    assert resp.status_code == 302


# ── POST /cruce ──


def test_post_sin_archivo_400(app_client):
    _login(app_client, ["cruce_produccion:write"])
    resp = app_client.post(
        "/cruce-produccion/cruce",
        headers={"X-Requested-With": "XMLHttpRequest"},
    )
    assert resp.status_code == 400
    data = resp.get_json()
    assert data["status"] == "error"
    assert data["status"] != "warning"
    assert isinstance(data["data"], dict)
    assert isinstance(data["errors"], list)


def test_post_cruce_faltantes_con_responsable(app_client, _mock_scan):
    _login(app_client, ["cruce_produccion:write"])
    resp = app_client.post(
        "/cruce-produccion/cruce",
        data={"file_upload": (io.BytesIO(_produccion_bytes()), "prod.xlsx")},
        content_type="multipart/form-data",
        headers={"X-Requested-With": "XMLHttpRequest"},
    )
    assert resp.status_code == 200
    data = resp.get_json()
    assert data["status"] == "success"
    assert data["status"] != "warning"
    payload = data["data"]
    assert payload["faltantes"] == [
        {
            "codigo": "FEV999",
            "numero_factura_original": "FEV999",
            "responsable": "Luis",
            "fec_factura": "2026-02-01",
            "estado_novedad": None,
            "facturador": "",
        }
    ]
    assert payload["revisar"] == []
    assert payload["resumen"]["total_faltantes"] == 1
    assert payload["resumen"]["total_produccion"] == 2
    # Contrato con el frontend (CruceResumen en page.tsx): keys exactas.
    assert set(payload["resumen"]) == {
        "total_produccion", "total_carpetas", "total_faltantes", "total_revisar",
        "total_desactualizados",
    }
    assert payload["desactualizados"] == []
    assert payload["resumen"]["total_desactualizados"] == 0
    assert payload["scanned_roots"] == ["//fake-root"]
    assert payload["last_scan_at"]
    assert payload["export_id"]

    # Round-trip: el export_id descarga el xlsx cacheado.
    _login(app_client, ["cruce_produccion"])
    with app_client.session_transaction() as sess:
        sess["_rate_limiter"] = []
    got = app_client.get(f"/cruce-produccion/export?id={payload['export_id']}")
    assert got.status_code == 200
    assert "spreadsheetml.sheet" in got.content_type


def test_post_sin_permiso_write_403(app_client):
    _login(app_client, ["cruce_produccion"])
    resp = app_client.post(
        "/cruce-produccion/cruce",
        data={"file_upload": (io.BytesIO(_produccion_bytes()), "prod.xlsx")},
        content_type="multipart/form-data",
        headers={"X-Requested-With": "XMLHttpRequest"},
    )
    assert resp.status_code == 403
    assert resp.get_json()["status"] == "error"


# ── GET /export ──


def test_export_id_invalido_400(app_client):
    _login(app_client, ["cruce_produccion"])
    resp = app_client.get(
        "/cruce-produccion/export?id=no-existe",
        headers={"X-Requested-With": "XMLHttpRequest"},
    )
    assert resp.status_code == 400
    data = resp.get_json()
    assert data["status"] == "error"
    assert data["status"] != "warning"
    assert isinstance(data["errors"], list)


def test_export_sin_id_400(app_client):
    _login(app_client, ["cruce_produccion"])
    resp = app_client.get(
        "/cruce-produccion/export",
        headers={"X-Requested-With": "XMLHttpRequest"},
    )
    assert resp.status_code == 400
    assert resp.get_json()["status"] == "error"


def test_post_cruce_incluye_desactualizados(app_client, monkeypatch):
    """Novedad pendiente sobre carpeta escaneada → sección desactualizados."""
    _login(app_client, ["cruce_produccion:write"])
    monkeypatch.setattr(routes, "get_roots", lambda: (["//fake-root"], "env", None))
    monkeypatch.setattr(routes, "detect_all", lambda _roots: _scan_result())
    monkeypatch.setattr(
        routes,
        "obtener_novedades",
        lambda: [{"factura": "FEV123", "estado": "S"}],
    )
    resp = app_client.post(
        "/cruce-produccion/cruce",
        data={"file_upload": (io.BytesIO(_produccion_bytes()), "prod.xlsx")},
        content_type="multipart/form-data",
        headers={"X-Requested-With": "XMLHttpRequest"},
    )
    assert resp.status_code == 200
    payload = resp.get_json()["data"]
    assert payload["resumen"]["total_desactualizados"] == 1
    assert payload["desactualizados"] == [
        {
            "codigo": "FEV123",
            "facturador": "Fact1",
            "full_path": "/root/Fact1/FEV123",
            "mtime_carpeta": None,
            "estado_novedad": "S",
            "motivo": "Novedad pendiente",
            "fecha_estado": None,
        }
    ]
    # Las filas revisar/faltantes no cambian salvo fec_factura.
    assert payload["resumen"]["total_faltantes"] == 1
    assert payload["revisar"] == []


def test_post_cruce_faltante_con_novedad_pendiente(app_client, monkeypatch):
    """Novedad S sobre código faltante → estado_novedad S + fec del Excel."""
    _login(app_client, ["cruce_produccion:write"])
    monkeypatch.setattr(routes, "get_roots", lambda: (["//fake-root"], "env", None))
    monkeypatch.setattr(routes, "detect_all", lambda _roots: _scan_result())
    monkeypatch.setattr(
        routes,
        "obtener_novedades",
        lambda: [{"factura": "FEV999", "estado": "S"}],
    )
    resp = app_client.post(
        "/cruce-produccion/cruce",
        data={"file_upload": (io.BytesIO(_produccion_bytes()), "prod.xlsx")},
        content_type="multipart/form-data",
        headers={"X-Requested-With": "XMLHttpRequest"},
    )
    assert resp.status_code == 200
    payload = resp.get_json()["data"]
    assert payload["faltantes"] == [
        {
            "codigo": "FEV999",
            "numero_factura_original": "FEV999",
            "responsable": "Luis",
            "fec_factura": "2026-02-01",
            "estado_novedad": "S",
            "facturador": "",
        }
    ]
