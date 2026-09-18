"""Tests for the procesar formatted export (strict TDD).

Covers tasks 1.1/1.2 + 2.1-2.3 + 3.1/3.2 + 4.2-4.4: constants, casing,
style/filename, full-vs-slice parity and the GET export contract.
"""

from __future__ import annotations

import re
from contextlib import ExitStack
from io import BytesIO
from pathlib import Path
from unittest.mock import patch

from openpyxl import load_workbook

from app import create_app
from app.constants.columnas import PROCESAR_CASING_MAP, PROCESAR_EXPORT_HEADERS
from app.services.procesar_export import (
    build_procesar_export_workbook,
    filename_procesar_export,
)
from app.utils.formatting import to_upper_safe

_APP = create_app()
_APP.config.update({"TESTING": True, "SECRET_KEY": "test-secret-key"})


def _row(**overrides):
    base = {
        "tipo_error": "Decimales",
        "tipo_factura": "Odontologia",
        "factura": "fac-001",
        "fec_factura": "2026-05-15",
        "responsable_cierra": "MARIA DEL PILAR MUÑOZ",
        "descripcion": "valor con decimales",
        "procedimiento": "C001 - Limpieza",
        "detalle": "revisar detalle",
        "regla": "#60",
    }
    base.update(overrides)
    return base


def _workbook(rows):
    buffer = build_procesar_export_workbook(rows)
    return load_workbook(BytesIO(buffer.read()))


class TestExportConstants:
    def test_headers_exact_order(self):
        assert PROCESAR_EXPORT_HEADERS == [
            "Fec. Factura",
            "Tipo de error",
            "Número Factura",
            "Regla",
            "Responsable Cierra",
            "Descripción",
            "Procedimiento",
            "Detalle",
        ]

    def test_casing_map_title_only_responsable(self):
        assert PROCESAR_CASING_MAP == {
            "Fec. Factura": "passthrough",
            "Tipo de error": "passthrough",
            "Número Factura": "upper",
            "Regla": "passthrough",
            "Responsable Cierra": "title",
            "Descripción": "upper",
            "Procedimiento": "passthrough",
            "Detalle": "upper",
        }


class TestToUpperSafe:
    def test_none_returns_empty(self):
        assert to_upper_safe(None) == ""

    def test_non_string_returns_empty(self):
        assert to_upper_safe(123) == ""
        assert to_upper_safe(["x"]) == ""

    def test_preserves_enie_and_tildes(self):
        assert to_upper_safe("muñoz") == "MUÑOZ"
        assert to_upper_safe("pérez") == "PÉREZ"

    def test_collapses_whitespace(self):
        assert to_upper_safe("  valor   con\nsaltos  ") == "VALOR CON SALTOS"


class TestCasingMatrix:
    def test_title_upper_passthrough(self):
        from datetime import datetime as _dt

        ws = _workbook([_row()]).active
        values = [c.value for c in ws[2]]
        assert values[0] == _dt(2026, 5, 15)  # Fec. Factura parsed
        assert values[1] == "Decimales"  # passthrough
        assert values[2] == "FAC-001"  # UPPER
        assert values[3] == "#60"  # passthrough
        assert values[4] == "Maria Del Pilar Muñoz"  # Title
        assert values[5] == "VALOR CON DECIMALES"  # UPPER
        assert values[6] == "C001 - Limpieza"  # passthrough
        assert values[7] == "REVISAR DETALLE"  # UPPER

    def test_enie_row(self):
        ws = _workbook([_row(responsable_cierra="JUAN MUÑOZ PÉREZ")]).active
        assert ws.cell(row=2, column=5).value == "Juan Muñoz Pérez"

    def test_nulls_yield_empty_cells(self):
        ws = _workbook([
            _row(factura=None, descripcion=None, detalle=None,
                 responsable_cierra=None, regla=None)
        ]).active
        assert ws.cell(row=2, column=3).value in (None, "")
        assert ws.cell(row=2, column=5).value in (None, "")
        assert ws.cell(row=2, column=6).value in (None, "")
        assert ws.cell(row=2, column=8).value in (None, "")

    def test_formula_injection_prefixed(self):
        ws = _workbook([_row(descripcion="=EVIL()", detalle="+cmd")]).active
        assert ws.cell(row=2, column=6).value == "'=EVIL()"
        assert ws.cell(row=2, column=8).value == "'+CMD"


class TestStyleAndFilename:
    def test_header_row_matches_novedades_tokens(self):
        ws = _workbook([_row()]).active
        assert [c.value for c in ws[1]] == PROCESAR_EXPORT_HEADERS
        assert ws["A1"].fill.fgColor.rgb == "001B5E20"
        assert ws["A1"].font.color.rgb == "00FFFFFF"
        assert ws["A1"].font.bold is True
        assert ws.cell(row=1, column=8).fill.fgColor.rgb == "001B5E20"

    def test_zebra_and_border(self):
        ws = _workbook([_row(), _row(factura="fac-002")]).active
        assert ws["A2"].fill.fgColor.rgb == "00E8F5E9"
        assert ws["A3"].fill.fgColor.rgb == "00FFFFFF"
        for cell in (ws["A2"], ws["H2"], ws["A3"], ws["A1"]):
            for side in ("left", "right", "top", "bottom"):
                assert getattr(cell.border, side).style == "thin"
                assert getattr(cell.border, side).color.rgb == "00A5D6A7"

    def test_widths_freeze_and_date_format(self):
        ws = _workbook([_row()]).active
        assert ws.column_dimensions["A"].width == 20
        assert ws.column_dimensions["H"].width == 20
        assert ws.freeze_panes == "A2"
        assert ws.cell(row=2, column=1).number_format == "DD/MM/YYYY"

    def test_headers_only_when_zero_rows(self):
        ws = _workbook([]).active
        assert ws.max_row == 1
        assert [c.value for c in ws[1]] == PROCESAR_EXPORT_HEADERS

    def test_filename_pattern(self):
        assert re.match(r"^procesar-[A-Z][a-z]{2}-\d{4}\.xlsx$",
                        filename_procesar_export())


# =============================================================================
# Route contract: POST export_id + thin GET export
# =============================================================================

def _login_admin(client):
    with client.session_transaction() as sess:
        sess["ce_authenticated"] = True
        sess["username"] = "admin"
        sess["rol"] = "admin"
        sess["permisos"] = ["*"]


def _normalized(i: int) -> dict:
    return {
        "tipo_error": "Decimales",
        "tipo_factura": "Odontologia",
        "factura": f"FAC-{i:04d}",
        "fec_factura": "2026-05-15",
        "responsable_cierra": "JUAN PEREZ",
        "descripcion": "valor con decimales",
        "procedimiento": "C001",
        "detalle": "revisar",
        "fecha_cierre_vacia": False,
        "regla": "",
    }


def _detect_result(rows: list[dict]):
    return {
        "status": "success",
        "data": {
            "problemas": {
                "problemas": {"normalizados": rows},
                "missing_columns": [],
                "tipos_procesados": ["Odontologia"],
            },
            "tipos_procesados": ["Odontologia"],
        },
        "errors": [],
    }


def _post_procesar(client, rows: list[dict]):
    with ExitStack() as stack:
        stack.enter_context(patch(
            "app.routes.procesar.save_temp_excel",
            return_value=(Path("dummy.xlsx"), None),
        ))
        stack.enter_context(patch(
            "app.routes.procesar.detect_problems_only",
            return_value=(_detect_result(rows), 200),
        ))
        return client.post(
            "/procesar/",
            data={"file_upload": (BytesIO(b"fake"), "input.xlsx")},
            content_type="multipart/form-data",
        )


class TestProcesarExportContract:
    def test_post_returns_export_id(self, app_client):
        _login_admin(app_client)
        resp = _post_procesar(app_client, [_normalized(1)])
        assert resp.status_code == 200
        data = resp.get_json()["data"]
        assert data["export_id"]

    def test_full_vs_slice_parity(self, app_client):
        _login_admin(app_client)
        resp = _post_procesar(app_client, [_normalized(i) for i in range(120)])
        export_id = resp.get_json()["data"]["export_id"]
        shown = sum(
            t["cantidad_mostradas"]
            for f in resp.get_json()["data"]["errores"]
            for t in f["tipos"]
        )
        assert shown == 50

        get_resp = app_client.get(f"/procesar/export?id={export_id}")
        assert get_resp.status_code == 200
        ws = load_workbook(BytesIO(get_resp.data)).active
        assert ws.max_row - 1 == 120

    def test_zero_errors_headers_only(self, app_client):
        _login_admin(app_client)
        resp = _post_procesar(app_client, [])
        export_id = resp.get_json()["data"]["export_id"]
        assert export_id

        get_resp = app_client.get(f"/procesar/export?id={export_id}")
        assert get_resp.status_code == 200
        ws = load_workbook(BytesIO(get_resp.data)).active
        assert ws.max_row == 1

    def test_unknown_id_400_no_bytes(self, app_client):
        _login_admin(app_client)
        resp = app_client.get("/procesar/export?id=AAAAAAAAAAAAAAAAAAAAAA")
        assert resp.status_code == 400
        assert resp.get_json()["status"] == "error"
        assert not resp.content_type.startswith(
            "application/vnd.openxmlformats")

    def test_missing_id_400(self, app_client):
        _login_admin(app_client)
        resp = app_client.get("/procesar/export")
        assert resp.status_code == 400
        assert resp.get_json()["status"] == "error"

    def test_expired_id_400(self, app_client):
        import json as _json
        import time as _time

        from app.utils import procesar_export_store as _store

        _login_admin(app_client)
        resp = _post_procesar(app_client, [_normalized(1)])
        export_id = resp.get_json()["data"]["export_id"]
        path = _store.temp_export_directory() / f"{export_id}.json"
        payload = _json.loads(path.read_text())
        payload["created"] = _time.time() - _store.PROCESAR_EXPORT_TTL_SECONDS - 1
        path.write_text(_json.dumps(payload))

        get_resp = app_client.get(f"/procesar/export?id={export_id}")
        assert get_resp.status_code == 400

    def test_over_cap_413(self, app_client, monkeypatch):
        from app.utils import procesar_export_store as _store

        _login_admin(app_client)
        monkeypatch.setattr(_store, "PROCESAR_EXPORT_MAX_ROWS", 5)
        resp = _post_procesar(app_client, [_normalized(i) for i in range(6)])
        export_id = resp.get_json()["data"]["export_id"]

        get_resp = app_client.get(f"/procesar/export?id={export_id}")
        assert get_resp.status_code == 413
        assert get_resp.get_json()["status"] == "error"

    def test_no_permission_403(self, fresh_client):
        with fresh_client.session_transaction() as sess:
            sess["ce_authenticated"] = True
            sess["username"] = "nobody"
            sess["rol"] = "usuario"
            sess["permisos"] = []
        resp = fresh_client.get(
            "/procesar/export?id=AAAAAAAAAAAAAAAAAAAAAA",
            headers={"X-Requested-With": "XMLHttpRequest"},
        )
        assert resp.status_code == 403
        assert resp.get_json()["status"] == "error"

    def test_post_disk_failure_still_success_without_id(self, app_client):
        _login_admin(app_client)
        with ExitStack() as stack:
            stack.enter_context(patch(
                "app.routes.procesar.save_temp_excel",
                return_value=(Path("dummy.xlsx"), None),
            ))
            stack.enter_context(patch(
                "app.routes.procesar.detect_problems_only",
                return_value=(_detect_result([_normalized(1)]), 200),
            ))
            stack.enter_context(patch(
                "app.routes.procesar.export_store.put",
                side_effect=OSError("disk full"),
            ))
            resp = app_client.post(
                "/procesar/",
                data={"file_upload": (BytesIO(b"fake"), "input.xlsx")},
                content_type="multipart/form-data",
            )
        assert resp.status_code == 200
        assert resp.get_json()["status"] == "success"
        assert "export_id" not in resp.get_json()["data"]
