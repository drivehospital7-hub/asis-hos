"""Tests for engine path in ambulatoria/detect_all.py.

NOTE: the legacy/OFF path was removed by design (is_rule_engine_enabled()
returns True hard-coded — "Engine always on"). Tests mocking engine OFF
(test_legacy_path_*) were deleted; only engine-ON tests remain.
"""

from __future__ import annotations

from unittest.mock import MagicMock, patch

import pytest
from openpyxl import Workbook

from app.constants import AREA_AMBULATORIA


def _regla(nombre: str, dominio: str, grupo: str | None):
    from app.models import Regla

    return Regla(
        id=abs(hash(nombre)) % 10_000 + 1, nombre=nombre, dominio=dominio,
        estado="active", version=1, prioridad=10, severidad="error",
        activo=True, grupo_error=grupo,
    )


class _RecordingDetector:
    """Fake RuleBasedDetector: records requested names, serves payloads."""

    instances: list[str] = []
    payloads: dict[str, list[dict]] = {}

    def __init__(self, name: str, session, **kwargs):
        type(self).instances.append(name)
        self._name = name

    def detect(self, *args, **kwargs):
        return [dict(p) for p in type(self).payloads.get(self._name, [])]

    @classmethod
    def reset(cls, payloads: dict | None = None) -> None:
        cls.instances = []
        cls.payloads = dict(payloads or {})


def _build_simple_sheet() -> tuple[Workbook, dict[str, int | None]]:
    """Build a workbook with minimal columns for testing."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Datos"

    headers = [
        "Número Factura", "Código", "Cantidad", "Vlr. Unitario",
        "Vlr. Procedimiento", "Tipo Doc.", "Edad", "Tipo Identificación",
        "Código Entidad Cobrar", "Entidad Afiliación", "Tipo Usuario",
        "Vlr. Copago", "Código CUPS", "Fec Factura", "Fecha Cierre",
        "Responsable Cierra",
    ]
    for col_idx, header in enumerate(headers, start=1):
        ws.cell(row=1, column=col_idx, value=header)

    ws.cell(row=2, column=1, value="FAC-001")

    indices = {h: i for i, h in enumerate(headers)}
    return wb, indices


class TestAmbulatoriaEngineToggle:
    """Tests for the engine toggle in detect_all_problems_ambulatoria."""

    def _make_mock_session(self) -> MagicMock:
        """Create a mock DB session that returns empty results."""
        session = MagicMock()
        mock_query = MagicMock()
        mock_query.filter.return_value = mock_query
        mock_query.order_by.return_value = mock_query
        mock_query.first.return_value = None
        mock_query.all.return_value = []
        session.query.return_value = mock_query
        return session

    @patch("app.database.get_session")
    @patch("app.services.engine.rule_based_detector.RuleBasedDetector")
    @patch("app.services.ambulatoria.detect_all.is_rule_engine_enabled", return_value=True)
    def test_engine_path_evaluates_resolver_rules_only(
        self, mock_enabled: MagicMock, mock_detector_cls: MagicMock,
        mock_get_session: MagicMock,
    ) -> None:
        """Dynamic discovery: only resolver-served rules evaluate (no fixed list)."""
        served = [
            _regla("valores_decimales", "transversal", "Decimales"),
            _regla("nueva_regla_ui", "ambulatoria", "Decimales"),
        ]
        fake_resolver = MagicMock()
        fake_resolver.resolve.return_value = served
        _RecordingDetector.reset({
            "valores_decimales": [{"factura": "FAC-001", "problema": "dec"}],
            "nueva_regla_ui": [{"factura": "FAC-001", "problema": "nuevo"}],
        })
        mock_get_session.return_value = MagicMock()
        mock_detector_cls.side_effect = _RecordingDetector

        from app.services.ambulatoria.detect_all import (
            detect_all_problems_ambulatoria,
        )
        with patch(
            "app.services.engine.domain_detection.RuleResolver",
            return_value=fake_resolver,
        ):
            wb, indices = _build_simple_sheet()
            result, responsables = detect_all_problems_ambulatoria(
                wb.active, indices,
            )

        assert set(_RecordingDetector.instances) == {
            "valores_decimales", "nueva_regla_ui",
        }
        assert "problemas" in result
        assert isinstance(result["problemas"], dict)
        assert "totales" in result
        assert result["area"] == AREA_AMBULATORIA
        assert len(result["problemas"]["decimales"]) == 2
        assert responsables == {}

    # NOTE: legacy/OFF path removed — engine is always on, so tests mocking
    # is_rule_engine_enabled → False were deleted (no OFF branch exists).

    @patch("app.database.get_session")
    @patch("app.services.engine.rule_based_detector.RuleBasedDetector")
    @patch("app.services.ambulatoria.detect_all.is_rule_engine_enabled", return_value=True)
    def test_engine_path_with_all_detectors(
        self, mock_enabled: MagicMock, mock_detector_cls: MagicMock,
        mock_get_session: MagicMock,
    ) -> None:
        """Engine path must produce problems dict with all keys present."""
        served = [
            _regla("valores_decimales", "transversal", "Decimales"),
            _regla("tipo_documento_edad_menor_7", "transversal", "Tipo Identificacion / Edad"),
            _regla("codigo_entidad", "transversal", "Codigo-Entidad-vs-Afiliacion"),
            _regla("tipo_usuario_valido", "transversal", "Tipo Usuario"),
            _regla("copago_entidad_valido", "transversal", "Copago vs Entidad"),
            _regla("cups_sin_contrato", "transversal", "Cups Sin Contrato"),
        ]
        fake_resolver = MagicMock()
        fake_resolver.resolve.return_value = served
        _RecordingDetector.reset()
        mock_get_session.return_value = MagicMock()
        mock_detector_cls.side_effect = _RecordingDetector

        from app.services.ambulatoria.detect_all import (
            detect_all_problems_ambulatoria,
        )
        with patch(
            "app.services.engine.domain_detection.RuleResolver",
            return_value=fake_resolver,
        ):
            wb, indices = _build_simple_sheet()
            result, _ = detect_all_problems_ambulatoria(wb.active, indices)

        assert set(_RecordingDetector.instances) == {r.nombre for r in served}

        problemas = result["problemas"]
        expected_keys = {
            "normalizados", "centros_de_costos", "ide_contrato",
            "cups_equivalentes", "decimales", "tipo_identificacion_edad",
            "tipo_identificacion_entidad", "codigo_entidad_vs_afiliacion",
            "tipo_usuario", "copago_entidad", "cups_sin_contrato",
        }
        for key in expected_keys:
            assert key in problemas, f"Missing key: {key}"
