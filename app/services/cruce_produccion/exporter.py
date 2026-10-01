"""Exporta el resultado del cruce a Excel + cache en disco.

Workbook con 3 hojas (Faltantes, Revisar, Soportes no actualizados) usando los estilos de
app/utils/formatting.py. El cache reusa app.utils.procesar_export_store
(put/get) sin copiar el módulo.
"""

from __future__ import annotations

import logging
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from app.utils import procesar_export_store as export_store
from app.utils.formatting import (
    auto_adjust_column_width,
    create_data_row_style,
    create_header_style,
)

logger = logging.getLogger(__name__)

FALTANTES_HEADERS: list[str] = [
    "Código",
    "Número Factura Original",
    "Responsable",
    "Fec. Factura",
    "Estado Novedad",
    "Facturador",
]

REVISAR_HEADERS: list[str] = [
    "Número Factura Original",
    "Responsable",
    "Fec. Factura",
]

FALTANTES_KEYS: list[str] = [
    "codigo",
    "numero_factura_original",
    "responsable",
    "fec_factura",
    "estado_novedad",
    "facturador",
]

REVISAR_KEYS: list[str] = [
    "numero_factura_original",
    "responsable",
    "fec_factura",
]

DESACTUALIZADOS_HEADERS: list[str] = [
    "Código",
    "Facturador",
    "Ruta carpeta",
    "Mtime carpeta",
    "Estado novedad",
    "Motivo",
    "Fecha estado",
]

DESACTUALIZADOS_KEYS: list[str] = [
    "codigo",
    "facturador",
    "full_path",
    "mtime_carpeta",
    "estado_novedad",
    "motivo",
    "fecha_estado",
]


def _write_sheet(
    ws: Worksheet,
    headers: list[str],
    keys: list[str],
    rows: list[dict[str, Any]],
) -> None:
    """Escribe headers con estilo + filas de datos con estilo."""
    header_style = create_header_style()
    for col, header in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = header_style["font"]
        cell.fill = header_style["fill"]
        cell.border = header_style["border"]
        cell.alignment = header_style["alignment"]
    data_style = create_data_row_style()
    for row_idx, item in enumerate(rows, start=2):
        for col, key in enumerate(keys, start=1):
            cell = ws.cell(row=row_idx, column=col, value=item.get(key))
            cell.fill = data_style["fill"]
            cell.border = data_style["border"]
            cell.alignment = data_style["alignment"]
    auto_adjust_column_width(ws)


def build_cruce_export_workbook(
    faltantes: list[dict[str, Any]],
    revisar: list[dict[str, Any]],
    desactualizados: list[dict[str, Any]] | None = None,
) -> BytesIO:
    """Construye el workbook del cruce en memoria."""
    rows_des = desactualizados if desactualizados is not None else []
    wb = Workbook()
    ws_falt = wb.active
    ws_falt.title = "Faltantes"
    _write_sheet(ws_falt, FALTANTES_HEADERS, FALTANTES_KEYS, faltantes)
    ws_rev = wb.create_sheet("Revisar")
    _write_sheet(ws_rev, REVISAR_HEADERS, REVISAR_KEYS, revisar)
    ws_des = wb.create_sheet("Soportes no actualizados")
    _write_sheet(ws_des, DESACTUALIZADOS_HEADERS, DESACTUALIZADOS_KEYS, rows_des)
    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    logger.info(
        "Workbook de cruce: %d faltantes, %d a revisar, %d desactualizados",
        len(faltantes),
        len(revisar),
        len(rows_des),
    )
    return buffer


def cache_cruce_result(
    faltantes: list[dict[str, Any]],
    revisar: list[dict[str, Any]],
    desactualizados: list[dict[str, Any]] | None = None,
) -> str:
    """Guarda el resultado en el export store y retorna el id opaco."""
    return export_store.put([{
        "faltantes": list(faltantes),
        "revisar": list(revisar),
        "desactualizados": list(desactualizados) if desactualizados is not None else [],
    }])


def get_cached_cruce_result(export_id: str) -> dict[str, Any] | None:
    """Recupera {"faltantes", "revisar", "desactualizados"} o None si falta/expiró."""
    rows = export_store.get(export_id)
    if not rows or not isinstance(rows[0], dict):
        return None
    payload = rows[0]
    if not isinstance(payload.get("faltantes"), list) or not isinstance(
        payload.get("revisar"), list
    ):
        return None
    des = payload.get("desactualizados", [])
    return {
        "faltantes": payload["faltantes"],
        "revisar": payload["revisar"],
        "desactualizados": des if isinstance(des, list) else [],
    }
